"""
Document generation skill.

Generates:
1. GST invoice PDF for a finalized bill.
2. Business analysis PPTX with real charts.
"""

from __future__ import annotations

import sqlite3
from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
    Paragraph,
)

from db.connection import get_connection


# Store generated artifacts here.
OUTPUT_DIR = Path(__file__).resolve().parent.parent / "assets"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def _get_bill(conn: sqlite3.Connection, bill_id: int):
    bill = conn.execute(
        """
        SELECT
            b.*,
            c.name AS customer_name
        FROM bills b
        LEFT JOIN customers c ON c.id = b.customer_id
        WHERE b.id = ?
        """,
        (bill_id,),
    ).fetchone()

    if bill is None:
        raise ValueError(f"No bill with id {bill_id}")

    return bill


def _get_bill_items(conn: sqlite3.Connection, bill_id: int):
    return conn.execute(
        """
        SELECT
            bi.*,
            p.name AS product_name,
            p.unit
        FROM bill_items bi
        JOIN products p ON p.id = bi.product_id
        WHERE bi.bill_id = ?
        ORDER BY bi.id
        """,
        (bill_id,),
    ).fetchall()


def generate_invoice_pdf(bill_id: int) -> dict:
    """
    Generate a GST invoice PDF for a finalized bill.

    The bill must already be finalized.
    """

    conn = get_connection()

    try:
        bill = _get_bill(conn, bill_id)

        if bill["status"] != "finalized":
            raise ValueError(
                f"Bill #{bill_id} is not finalized. "
                "Only finalized bills can produce invoices."
            )

        items = _get_bill_items(conn, bill_id)

        if not items:
            raise ValueError(f"Bill #{bill_id} has no items.")

        filename = f"invoice_bill_{bill_id}.pdf"
        output_path = OUTPUT_DIR / filename

        shop_name = "My Kirana Store"

        # Try to get configured shop name without making this
        # document skill dependent on bot/config internals.
        try:
            from bot import config

            if config.SHOP_NAME:
                shop_name = config.SHOP_NAME
            shop_gstin = config.SHOP_GSTIN or ""
        except Exception:
            shop_gstin = ""

        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            rightMargin=15 * mm,
            leftMargin=15 * mm,
            topMargin=15 * mm,
            bottomMargin=15 * mm,
        )

        styles = getSampleStyleSheet()

        title_style = styles["Title"]
        title_style.fontSize = 20
        title_style.leading = 24

        normal = styles["Normal"]
        normal.fontSize = 9
        normal.leading = 12

        small = styles["Normal"]
        small.fontSize = 8
        small.leading = 10

        story = []

        story.append(Paragraph(shop_name, title_style))
        story.append(Paragraph("GST TAX INVOICE", styles["Heading2"]))

        if shop_gstin:
            story.append(
                Paragraph(
                    f"<b>GSTIN:</b> {shop_gstin}",
                    normal,
                )
            )

        created_at = bill["finalized_at"] or bill["created_at"]

        story.append(
            Paragraph(
                f"<b>Invoice No:</b> INV-{bill_id:06d}<br/>"
                f"<b>Date:</b> {created_at}<br/>"
                f"<b>Payment:</b> {bill['payment_mode'] or '-'}"
                + (
                    f"<br/><b>Payment Ref:</b> {bill['payment_ref']}"
                    if bill["payment_ref"]
                    else ""
                )
                + (
                    f"<br/><b>Customer:</b> {bill['customer_name']}"
                    if bill["customer_name"]
                    else ""
                ),
                normal,
            )
        )

        story.append(Spacer(1, 8 * mm))

        table_data = [
            [
                "Item",
                "Qty",
                "Unit Price",
                "GST %",
                "Taxable",
                "CGST",
                "SGST",
                "Total",
            ]
        ]

        for item in items:
            line_total = float(item["line_total"])
            gst_rate = float(item["gst_rate"])

            # sell_price is GST-inclusive.
            taxable = (
                line_total / (1 + gst_rate / 100)
                if gst_rate
                else line_total
            )

            total_gst = line_total - taxable
            cgst = total_gst / 2
            sgst = total_gst / 2

            table_data.append(
                [
                    str(item["product_name"]),
                    f"{float(item['quantity']):g} {item['unit']}",
                    f"₹{float(item['unit_price']):.2f}",
                    f"{gst_rate:g}%",
                    f"₹{taxable:.2f}",
                    f"₹{cgst:.2f}",
                    f"₹{sgst:.2f}",
                    f"₹{line_total:.2f}",
                ]
            )

        table = Table(
            table_data,
            repeatRows=1,
            colWidths=[
                38 * mm,
                23 * mm,
                25 * mm,
                17 * mm,
                25 * mm,
                21 * mm,
                21 * mm,
                25 * mm,
            ],
        )

        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )

        story.append(table)
        story.append(Spacer(1, 8 * mm))

        summary_data = [
            ["Subtotal", f"₹{float(bill['subtotal'] or 0):.2f}"],
            ["CGST", f"₹{float(bill['cgst_total'] or 0):.2f}"],
            ["SGST", f"₹{float(bill['sgst_total'] or 0):.2f}"],
            [
                "Grand Total",
                f"₹{float(bill['grand_total'] or 0):.2f}",
            ],
        ]

        summary = Table(
            summary_data,
            colWidths=[45 * mm, 35 * mm],
            hAlign="RIGHT",
        )

        summary.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                    ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                    ("BACKGROUND", (0, -1), (-1, -1), colors.lightgrey),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )

        story.append(summary)
        story.append(Spacer(1, 10 * mm))

        story.append(
            Paragraph(
                "Thank you for shopping with us.",
                normal,
            )
        )

        doc.build(story)

        return {
            "success": True,
            "type": "invoice_pdf",
            "bill_id": bill_id,
            "filename": filename,
            "path": str(output_path.resolve()),
        }

    finally:
        conn.close()


def _sales_data(start_date: str, end_date: str):
    conn = get_connection()

    try:
        sales = conn.execute(
            """
            SELECT
                COUNT(*) AS bill_count,
                COALESCE(SUM(subtotal), 0) AS subtotal,
                COALESCE(SUM(cgst_total), 0) AS cgst,
                COALESCE(SUM(sgst_total), 0) AS sgst,
                COALESCE(SUM(grand_total), 0) AS revenue
            FROM bills
            WHERE status = 'finalized'
              AND date(finalized_at) BETWEEN date(?) AND date(?)
            """,
            (start_date, end_date),
        ).fetchone()

        top_items = conn.execute(
            """
            SELECT
                p.name,
                SUM(bi.quantity) AS quantity,
                SUM(bi.line_total) AS revenue
            FROM bill_items bi
            JOIN bills b ON b.id = bi.bill_id
            JOIN products p ON p.id = bi.product_id
            WHERE b.status = 'finalized'
              AND date(b.finalized_at) BETWEEN date(?) AND date(?)
            GROUP BY p.id, p.name
            ORDER BY revenue DESC
            LIMIT 10
            """,
            (start_date, end_date),
        ).fetchall()

        stock = conn.execute(
            """
            SELECT
                COUNT(*) AS total_products,
                SUM(
                    CASE
                        WHEN quantity <= reorder_level THEN 1
                        ELSE 0
                    END
                ) AS low_stock
            FROM products
            """
        ).fetchone()

        gst = conn.execute(
            """
            SELECT
                gst_rate,
                COALESCE(SUM(
                    line_total - (
                        line_total / (1 + gst_rate / 100)
                    )
                ), 0) AS gst_amount
            FROM bill_items bi
            JOIN bills b ON b.id = bi.bill_id
            WHERE b.status = 'finalized'
              AND date(b.finalized_at) BETWEEN date(?) AND date(?)
            GROUP BY gst_rate
            ORDER BY gst_rate
            """,
            (start_date, end_date),
        ).fetchall()

        payments = conn.execute(
            """
            SELECT
                payment_mode,
                COUNT(*) AS bill_count,
                COALESCE(SUM(grand_total), 0) AS amount
            FROM bills
            WHERE status = 'finalized'
              AND date(finalized_at) BETWEEN date(?) AND date(?)
            GROUP BY payment_mode
            ORDER BY amount DESC
            """,
            (start_date, end_date),
        ).fetchall()

        return sales, top_items, stock, gst, payments

    finally:
        conn.close()


def generate_analysis_pptx(
    start_date: str,
    end_date: str,
) -> dict:
    """
    Generate a business analysis PowerPoint with real charts.
    """

    sales, top_items, stock, gst, payments = _sales_data(
        start_date,
        end_date,
    )

    filename = f"sales_analysis_{start_date}_to_{end_date}.pptx"
    output_path = OUTPUT_DIR / filename

    chart_paths = []

    # ------------------------------------------------------------
    # Chart 1: Top items
    # ------------------------------------------------------------

    if top_items:
        names = [str(row["name"]) for row in top_items]
        revenues = [float(row["revenue"] or 0) for row in top_items]

        plt.figure(figsize=(10, 5))
        plt.bar(names, revenues)
        plt.title("Top Items by Revenue")
        plt.ylabel("Revenue (INR)")
        plt.xticks(rotation=35, ha="right")
        plt.tight_layout()

        top_items_chart = OUTPUT_DIR / "_chart_top_items.png"
        plt.savefig(top_items_chart, dpi=150)
        plt.close()

        chart_paths.append(top_items_chart)

    # ------------------------------------------------------------
    # Chart 2: Payment split
    # ------------------------------------------------------------

    if payments:
        labels = [str(row["payment_mode"]) for row in payments]
        amounts = [float(row["amount"] or 0) for row in payments]

        plt.figure(figsize=(7, 5))
        plt.pie(
            amounts,
            labels=labels,
            autopct="%1.1f%%",
        )
        plt.title("Payment Mode Split")

        payment_chart = OUTPUT_DIR / "_chart_payments.png"
        plt.savefig(payment_chart, dpi=150)
        plt.close()

        chart_paths.append(payment_chart)

    # ------------------------------------------------------------
    # Chart 3: GST slabs
    # ------------------------------------------------------------

    if gst:
        labels = [f"{float(row['gst_rate']):g}%" for row in gst]
        amounts = [float(row["gst_amount"] or 0) for row in gst]

        plt.figure(figsize=(8, 5))
        plt.bar(labels, amounts)
        plt.title("GST Collected by Slab")
        plt.xlabel("GST Rate")
        plt.ylabel("GST Amount (INR)")
        plt.tight_layout()

        gst_chart = OUTPUT_DIR / "_chart_gst.png"
        plt.savefig(gst_chart, dpi=150)
        plt.close()

        chart_paths.append(gst_chart)

    # ------------------------------------------------------------
    # Build PPTX
    # ------------------------------------------------------------

    prs = Presentation()

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Supermarket Business Analysis"

    subtitle = slide.placeholders[1]
    subtitle.text = f"{start_date} to {end_date}"

    # Summary slide
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Business Summary"

    text_frame = slide.placeholders[1].text_frame
    text_frame.clear()

    summary_lines = [
        f"Total finalized bills: {int(sales['bill_count'] or 0)}",
        f"Revenue: ₹{float(sales['revenue'] or 0):,.2f}",
        f"Subtotal: ₹{float(sales['subtotal'] or 0):,.2f}",
        f"CGST collected: ₹{float(sales['cgst'] or 0):,.2f}",
        f"SGST collected: ₹{float(sales['sgst'] or 0):,.2f}",
        f"Low-stock products: {int(stock['low_stock'] or 0)}",
        f"Total products: {int(stock['total_products'] or 0)}",
    ]

    for index, line in enumerate(summary_lines):
        if index == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()

        p.text = line
        p.font.size = Pt(20)

    # Top-items chart
    if top_items:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Top Items by Revenue"
        slide.shapes.add_picture(
            str(top_items_chart),
            Inches(0.7),
            Inches(1.4),
            width=Inches(8.5),
        )

    # Payment chart
    if payments:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Payment Mode Split"
        slide.shapes.add_picture(
            str(payment_chart),
            Inches(2.0),
            Inches(1.3),
            width=Inches(6.5),
        )

    # GST chart
    if gst:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "GST Collected by Slab"
        slide.shapes.add_picture(
            str(gst_chart),
            Inches(1.0),
            Inches(1.3),
            width=Inches(8.0),
        )

    # Insights
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Key Insights"

    tf = slide.placeholders[1].text_frame
    tf.clear()

    insights = []

    if top_items:
        best = top_items[0]
        insights.append(
            f"Top revenue item: {best['name']} "
            f"(₹{float(best['revenue'] or 0):,.2f})"
        )

    if payments:
        dominant = payments[0]
        insights.append(
            f"Most-used payment mode by value: "
            f"{dominant['payment_mode']} "
            f"(₹{float(dominant['amount'] or 0):,.2f})"
        )

    insights.append(
        f"{int(stock['low_stock'] or 0)} product(s) are at or below "
        "their reorder level."
    )

    if not insights:
        insights.append("No finalized sales were found for this period.")

    for index, insight in enumerate(insights):
        if index == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()

        p.text = insight
        p.font.size = Pt(20)

    prs.save(output_path)

    # Remove temporary chart files after embedding them.
    for chart_path in chart_paths:
        try:
            chart_path.unlink()
        except OSError:
            pass

    return {
        "success": True,
        "type": "analysis_pptx",
        "start_date": start_date,
        "end_date": end_date,
        "filename": filename,
        "path": str(output_path.resolve()),
    }